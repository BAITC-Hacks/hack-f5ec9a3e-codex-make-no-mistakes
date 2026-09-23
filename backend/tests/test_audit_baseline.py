"""Regression checks against frozen source bytes; never rewrite a changed baseline automatically."""

import hashlib
import json
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
BASELINE = json.loads((ROOT / "docs/sources/audit-baseline.json").read_text(encoding="utf-8"))
INVENTORY = json.loads((ROOT / "docs/sources/workbook-inventory.json").read_text(encoding="utf-8"))


def test_original_bytes_and_every_sheet_are_locked():
    manifest = json.loads((ROOT / "docs/sources/manifest.json").read_text(encoding="utf-8"))
    retained = [item for item in manifest if "path" in item]
    assert len(retained) == 14
    for item in retained:
        data = (ROOT / item["path"]).read_bytes()
        assert len(data) == item["bytes"]
        assert hashlib.sha256(data).hexdigest() == item["sha256"], item["path"]
    assert len(INVENTORY) == BASELINE["workbooks"]
    count = 0
    for entry in INVENTORY:
        workbook = openpyxl.load_workbook(ROOT / entry["path"], read_only=True)
        assert workbook.sheetnames == [s["name"] for s in entry["sheets"]]
        count += len(workbook.sheetnames)
        workbook.close()
    assert count == BASELINE["sheets"]


def test_sales_signs_missing_quantities_and_scope_are_locked():
    for supplier, expected in BASELINE["movements"].items():
        path = next((ROOT / "docs/data" / supplier).glob("Динамика*.xlsx"))
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        counts = Counter()
        skus, warehouses = set(), set()
        for row in workbook.active.iter_rows(min_row=2, max_col=8, values_only=True):
            if not row[3]:  # Footer is not a SKU movement.
                continue
            counts["rows"] += 1
            skus.add(row[3])
            warehouses.add(row[6])
            if row[7] is None:
                counts["missing_quantity"] += 1
            else:
                counts["positive" if row[7] > 0 else "negative" if row[7] < 0 else "zero"] += 1
        counts["unique_skus"] = len(skus)
        assert dict(counts) == expected
        assert sorted(warehouses) == BASELINE["warehouse_names"]
        workbook.close()


def test_systeme_stock_and_formula_defect_are_locked():
    path = next((ROOT / "docs/data/Systeme electric").glob("Товар*.xlsx"))
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook["TDSheet"].iter_rows(min_row=3, max_row=499, values_only=True))
    expected = BASELINE["systeme_report"]
    assert len(rows) == expected["sku_rows"]
    for row in rows:
        assert row[49] - row[50] == row[51]
    for key, column in (("on_hand", 49), ("reserved", 50), ("free", 51), ("incoming", 54)):
        assert sum(row[column] for row in rows) == expected[key]
    demo = expected["demo"]
    row = rows[demo["row"] - 3]
    assert row[2] == demo["sku"] and row[1] == demo["article"]
    assert [row[i] for i in (49, 50, 51, 54)] == [
        demo[k] for k in ("on_hand", "reserved", "free", "incoming")
    ]
    workbook.close()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    assert workbook["TDSheet"]["AP3"].value == expected["thirteen_month_formula"]
    assert workbook["TDSheet"]["AQ3"].value == expected["average_formula"]
    workbook.close()


def test_iek_errors_and_duplicate_identity_are_locked():
    path = next((ROOT / "docs/data/IEK").glob("MOQ*.xlsx"))
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
    expected = BASELINE["iek_moq"]
    assert sum(row[4] == "#N/A" for row in rows) == expected["error_cells"]
    assert [i + 2 for i, row in enumerate(rows) if row[1] == expected["duplicate_code"]] == expected[
        "duplicate_rows"
    ]
    workbook.close()
