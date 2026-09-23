import hashlib
import os
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

import openpyxl
import pytest
from alembic import command
from alembic.config import Config
from sqlalchemy import create_engine, inspect, text
from sqlalchemy.engine import make_url

from replenishment.cli.import_excel import import_workbook, supplier_for
from replenishment.intake.adapters.normalization import classify, number, shipment_header
from replenishment.intake.adapters.xlsx import Workbook


def source(tmp_path, *, broken=False):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Movements"
    sheet.append(["Дата", "Номер", "Документ", "Код", "Наименование", "Ед.", "Склад", "Количество"])
    sku = "000_" + uuid4().hex
    sheet.append(["20.07.2023 16:15:38", "001", "Return", sku, "Product", "шт", "Алматы", -10])
    sheet.append(["21.07.2023 16:15:38", "002", "Sale", sku, "Product", "шт", "Алматы", None])
    sheet.append(["22.07.2023 16:15:38", "003", "Sale", sku, "Product", "шт", "Алматы", 0])
    sheet.append(
        [
            "bad date" if broken else "23.07.2023 16:15:38",
            "004",
            "Sale",
            sku,
            "Product",
            "шт",
            "Алматы",
            "#N/A",
        ]
    )
    sheet.merge_cells("J1:K1")
    sheet["J1"] = "Extra raw field"
    sheet["K2"] = "=1/3"
    sheet["L2"] = 1.1
    sheet.row_dimensions[3].hidden = True
    path = tmp_path / f"IEK_{uuid4().hex}.xlsx"
    book.save(path)
    return path, sku


def test_raw_lexemes_formula_and_layout(tmp_path):
    path, _ = source(tmp_path)
    target = BytesIO()
    with ZipFile(path) as archive, ZipFile(target, "w") as output:
        for member in archive.namelist():
            data = archive.read(member)
            if member == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<v>1.1</v>", b"<v>1.10000000000000000001</v>")
                data = data.replace(b"<f>1/3</f><v />", b"<f>1/3</f><v>0.333333333333333333</v>")
            output.writestr(member, data)
    book = Workbook(target.getvalue())
    try:
        sheet = book.sheets[0]
        rows = dict(book.rows(sheet))
        assert sheet.layout["merged_ranges"] == ["J1:K1"]
        assert any(row.get("hidden") == "1" for row in sheet.layout["special_rows"])
        assert rows[2]["L"]["value"] == "1.10000000000000000001"
        assert rows[2]["K"]["formula"] == "=1/3"
        assert rows[2]["K"]["cached_value"] == "0.333333333333333333"
        assert number(rows[2], "K") == Decimal("0.333333333333")
        assert "H" not in rows[3]
        assert rows[4]["H"]["value"] == "0"
        assert rows[5]["H"]["type"] == "e"
        assert number(rows[5], "H") is None
    finally:
        book.close()


def test_supplier_cannot_be_guessed_for_ambiguous_file():
    with pytest.raises(ValueError, match="Cannot infer"):
        supplier_for(Path("sales.xlsx"))
    assert supplier_for(Path("sales.xlsx"), "systeme") == "systeme"
    with pytest.raises(ValueError, match="conflicts"):
        supplier_for(Path("IEK/sales.xlsx"), "systeme")


def test_all_supplied_templates_and_shipment_dates():
    root = Path(__file__).resolve().parents[2] / "docs/data"
    templates = []
    from itertools import islice

    for path in root.rglob("*.xlsx"):
        workbook = Workbook(path.read_bytes())
        try:
            for sheet in workbook.sheets:
                headers = dict(islice(workbook.rows(sheet), 3))
                templates.append(classify(headers, supplier_for(path))[0])
                if templates[-1] == "shipments":
                    assert (
                        shipment_header(headers[1]["D"]["value"])["expected_on"].isoformat() == "2026-10-10"
                    )
        finally:
            workbook.close()
    assert len(templates) == 14
    assert set(templates) == {
        "moq",
        "movements",
        "monthly_stock",
        "monthly_sales",
        "shipments",
        "snapshot",
        "seasonality",
    }


@pytest.mark.postgres
def test_real_snapshot_provenance_and_ambiguities(import_engine):
    root = Path(__file__).resolve().parents[2] / "docs/data/Systeme electric"
    path = next(root.glob("Товар*.xlsx"))
    import_workbook(import_engine, path)
    with import_engine.connect() as connection:
        stock = connection.execute(
            text("""
            SELECT s.metric,s.quantity,s.warehouse_id,s.date_basis
            FROM inventory_observations s JOIN catalog_products p ON p.id=s.product_id
            WHERE p.sku='300200745_' AND s.normalizer_version='v1'
        """)
        ).all()
        assert {r.metric: r.quantity for r in stock}["free"] == Decimal(23)
        assert all(r.warehouse_id is None and r.date_basis == "filename" for r in stock)
        shipment = connection.execute(
            text("""
            SELECT s.expected_on,s.date_basis,l.quantity
            FROM supply_shipments s JOIN supply_shipment_lines l ON l.shipment_id=s.id
            JOIN catalog_products p ON p.id=l.product_id WHERE p.sku='300200745_'
        """)
        ).one()
        assert tuple(shipment) == (None, "unknown", Decimal(120))
        assert (
            connection.scalar(
                text("""
            SELECT count(*) FROM demand_monthly_sales s JOIN catalog_products p ON p.id=s.product_id
            WHERE p.sku='300200745_'
        """)
            )
            == 33
        )
        assert (
            connection.scalar(
                text("""
            SELECT count(*) FROM intake_findings WHERE code='thirteen_month_total'
        """)
            )
            == 1
        )


@pytest.fixture(scope="module")
def import_engine():
    url = os.environ.get("IMPORT_TEST_DATABASE_URL")
    if not url:
        pytest.skip("Set IMPORT_TEST_DATABASE_URL to a dedicated *_import_test PostgreSQL database")
    assert make_url(url).database.endswith("_import_test")
    engine = create_engine(url)
    previous = os.environ.get("DATABASE_URL")
    try:
        if not inspect(engine).has_table("alembic_version"):
            os.environ["DATABASE_URL"] = url
            command.upgrade(Config(str(Path(__file__).resolve().parents[1] / "alembic.ini")), "head")
        yield engine
    finally:
        engine.dispose()
        if previous is None:
            os.environ.pop("DATABASE_URL", None)
        else:
            os.environ["DATABASE_URL"] = previous


@pytest.mark.postgres
def test_atomic_idempotent_import_and_new_normalizer(import_engine, tmp_path):
    path, sku = source(tmp_path)
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: import_workbook(import_engine, path), range(2)))
    assert sorted(r["status"] for r in results) == ["imported", "skipped"]
    with import_engine.connect() as connection:
        quantities = (
            connection.execute(
                text("""
            SELECT m.quantity FROM demand_movements m JOIN catalog_products p ON p.id=m.product_id
            WHERE p.sku=:sku ORDER BY occurred_at
        """),
                {"sku": sku},
            )
            .scalars()
            .all()
        )
        assert quantities == [Decimal(-10), None, Decimal(0), None]
        raw_before = connection.scalar(text("SELECT count(*) FROM intake_rows"))
    result = import_workbook(import_engine, path, version="v1-replay-test")
    assert result["status"] == "imported" and result["counts"]["movements"] == 4
    with import_engine.connect() as connection:
        assert connection.scalar(text("SELECT count(*) FROM intake_rows")) == raw_before
        assert (
            connection.scalar(
                text("""
            SELECT count(*) FROM demand_movements m
            JOIN catalog_products p ON p.id=m.product_id WHERE p.sku=:sku
        """),
                {"sku": sku},
            )
            == 8
        )
        assert (
            connection.scalar(
                text("SELECT content FROM intake_workbooks WHERE sha256=:digest"),
                {"digest": hashlib.sha256(path.read_bytes()).hexdigest()},
            )
            == path.read_bytes()
        )
    with pytest.raises(ValueError, match="another supplier|conflicts"):
        import_workbook(import_engine, path, supplier="systeme")


@pytest.mark.postgres
def test_failed_workbook_leaves_no_partial_records(import_engine, tmp_path):
    path, sku = source(tmp_path, broken=True)
    with pytest.raises(ValueError):
        import_workbook(import_engine, path)
    with import_engine.connect() as connection:
        assert (
            connection.scalar(text("SELECT count(*) FROM catalog_products WHERE sku=:sku"), {"sku": sku}) == 0
        )
        assert (
            connection.scalar(
                text("SELECT count(*) FROM intake_workbooks WHERE sha256=:digest"),
                {"digest": hashlib.sha256(path.read_bytes()).hexdigest()},
            )
            == 0
        )
