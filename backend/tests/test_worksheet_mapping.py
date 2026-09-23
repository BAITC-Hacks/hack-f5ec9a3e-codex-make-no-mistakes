from copy import deepcopy
from decimal import Decimal
from itertools import islice
from pathlib import Path

import pytest
from openpyxl.utils.cell import column_index_from_string, get_column_letter

from replenishment.cli.import_excel import FINDINGS
from replenishment.intake.adapters.normalization import (
    classify,
    mapping_from_payload,
    month_header,
    number,
    observations,
    value,
)
from replenishment.intake.adapters.xlsx import Workbook

ROOT = Path(__file__).resolve().parents[2]
MONTHLY_SALES = ROOT / "docs/data/IEK/Ежемесячные продажи в количественном выражении за последние 2 года.xlsx"
MOVEMENTS = ROOT / "docs/data/IEK/Динамика продаж_2025-2026.xlsx"


def _shift(cells, offset=2):
    return {
        get_column_letter(column_index_from_string(column) + offset): deepcopy(cell)
        for column, cell in cells.items()
    }


def _source_rows(path, count=4):
    workbook = Workbook(path.read_bytes())
    sheet = workbook.sheets[0]
    rows = list(islice(workbook.rows(sheet), count))
    workbook.close()
    return rows


def test_mapping_follows_displaced_columns_and_header_rows_from_real_sales_source():
    rows = _source_rows(MONTHLY_SALES)
    original_headers = dict(rows[:2])
    original_row, original_cells = rows[2]
    original = classify(original_headers, "iek")
    shifted_headers = {row: _shift(cells) for row, cells in original_headers.items()}
    shifted_cells = _shift(original_cells)
    shifted = classify(shifted_headers, "iek")

    old_month = next(
        (column, period)
        for column, period in original.month_columns
        if number(original_cells, column) is not None
    )
    new_month = next(period for column, period in shifted.month_columns if period == old_month[1])
    new_month_column = next(column for column, period in shifted.month_columns if period == new_month)
    assert shifted.kind == "monthly_sales"
    assert value(shifted_cells, shifted.sku_col) == value(original_cells, original.sku_col)
    assert number(shifted_cells, new_month_column) == number(original_cells, old_month[0])

    multiline_headers = {row + 1: deepcopy(cells) for row, cells in original_headers.items()}
    multiline = classify(multiline_headers, "iek")
    assert multiline.header_rows == (2, 3)
    assert multiline.month_columns[0][1] == month_header(
        value(original_headers[1], original.month_columns[0][0])
    )


def test_invalid_llm_cell_citation_cannot_become_normalized_observation():
    rows = _source_rows(MONTHLY_SALES)
    headers = dict(rows[:2])
    mapping = classify(headers, "iek")
    month_column, period = mapping.month_columns[0]
    payload = {
        "report_type": "monthly_sales",
        "header_row": 1,
        "data_start_row": 3,
        "columns": {
            "sku": mapping.sku_col,
            "months": [{"column": month_column, "period": period.isoformat()}],
        },
        "source_cells": [
            {"cell": "ZZ999", "purpose": "sku header"},
            {"cell": f"{mapping.sku_col}1", "purpose": "sku header"},
            {"cell": f"{month_column}1", "purpose": "month header"},
        ],
    }
    with pytest.raises(ValueError, match="source cell"):
        mapping_from_payload(payload, headers, "iek")
    payload["source_cells"] = payload["source_cells"][1:]
    payload["columns"]["months"][0]["period"] = "2099-01-01"
    with pytest.raises(ValueError, match="period"):
        mapping_from_payload(payload, headers, "iek")


def test_blank_zero_error_and_missing_formula_cache_remain_distinct_on_real_row():
    rows = _source_rows(MONTHLY_SALES)
    headers = dict(rows[:2])
    row, source_cells = rows[2]
    mapping = classify(headers, "iek")
    month_column, _ = mapping.month_columns[0]

    blank = deepcopy(source_cells)
    blank.pop(month_column, None)
    zero = deepcopy(source_cells)
    zero[month_column] = {"type": "n", "value": "0", "number_format": "General"}
    error = deepcopy(source_cells)
    error[month_column] = {"type": "e", "value": "#N/A", "number_format": "General"}
    formula = deepcopy(source_cells)
    formula[month_column] = {
        "type": "f",
        "formula": "=1/3",
        "cached_type": "n",
        "cached_value": None,
        "number_format": "General",
    }

    assert value(blank, month_column) is None and number(blank, month_column) is None
    assert value(zero, month_column) == "0" and number(zero, month_column) == Decimal("0.000000000000")
    assert value(error, month_column) is None and number(error, month_column) is None
    assert value(formula, month_column) is None and number(formula, month_column) is None
    assert formula[month_column]["formula"] == "=1/3"
    assert row > mapping.header_end


def test_systeme_multiple_uses_cited_header_above_quantity_subheader():
    path = (
        ROOT
        / "docs/data/Systeme electric/Ежемесячные продажи в кол-м выражении SystemElectric 2024-2026.xlsx"
    )
    rows = _source_rows(path)
    headers = dict(rows[:2])
    mapping = classify(headers, "systeme")
    row, cells = rows[2]
    rule = next(
        item
        for table, item in observations(mapping.kind, "systeme", cells, row, headers, None, None, mapping)
        if table == "quantity_rules"
    )
    assert rule["interpretation_confirmed"] is True
    assert "кратность" in rule["label"].casefold()


def test_missing_movement_warehouse_stays_unknown():
    rows = _source_rows(MOVEMENTS, 3)
    headers = {rows[0][0]: rows[0][1]}
    mapping = classify(headers, "iek")
    row, cells = rows[1]
    missing = deepcopy(cells)
    missing.pop(mapping.fields["warehouse"], None)
    movement = next(observations(mapping.kind, "iek", missing, row, headers, "шт", None, mapping=mapping))
    assert movement[1]["warehouse_name"] is None
    assert "missing_movement_warehouse" in FINDINGS


def test_supplied_stock_and_snapshot_sources_keep_observation_domains_separate():
    stock_path = ROOT / "docs/data/Systeme electric/Ежемесячные остатки SystemElectric 2024-2026.xlsx"
    snapshot_path = ROOT / "docs/data/Systeme electric/Товар в пути_SystemElectric на 22.09.2026.xlsx"
    workbook = Workbook(stock_path.read_bytes())
    try:
        sheet = workbook.sheets[0]
        rows = list(islice(workbook.rows(sheet), 5))
        headers = dict(rows[:1])
        mapping = classify(headers, "systeme")
        data_row, cells = rows[1]
        tables = list(
            observations(
                mapping.kind,
                "systeme",
                cells,
                data_row,
                headers,
                value(cells, mapping.unit_col),
                None,
                mapping=mapping,
            )
        )
        assert mapping.kind == "monthly_stock"
        assert {table for table, _ in tables} == {"monthly_stock"}
    finally:
        workbook.close()

    workbook = Workbook(snapshot_path.read_bytes())
    try:
        sheet = workbook.sheets[0]
        rows = list(islice(workbook.rows(sheet), 4))
        headers = dict(rows[:3])
        mapping = classify(headers, "systeme")
        data_row, cells = rows[2]
        tables = list(
            observations(mapping.kind, "systeme", cells, data_row, headers, None, None, mapping=mapping)
        )
        inventory_metrics = {
            "showroom",
            "trading_area",
            "distribution_center",
            "retail",
            "on_hand",
            "reserved",
            "free",
        }
        assert mapping.kind == "snapshot"
        assert all(
            item.get("metric") not in inventory_metrics for table, item in tables if table == "report_metrics"
        )
        assert {table for table, _ in tables} == {
            "monthly_sales",
            "report_metrics",
            "shipment_lines",
            "stock",
        }
    finally:
        workbook.close()
