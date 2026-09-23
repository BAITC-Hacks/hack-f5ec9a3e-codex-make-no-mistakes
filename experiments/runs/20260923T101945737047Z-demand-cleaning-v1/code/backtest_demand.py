"""Read-only, rolling historical forecast experiment; see docs/DEMAND_EXPERIMENT.md.

Run: python scripts/backtest_demand.py
Dependencies: numpy, openpyxl. Source workbooks are never modified.
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts/demand-experiment"
START = date(2025, 1, 1)
END = date(2026, 8, 31)
MODELS = ["sultan_mean59", "mean28", "weekly_ewma", "robust_ewma",
          "seasonal364", "seasonal_blend"]
SHORT_MODELS = MODELS[:4]


def forecast(history: np.ndarray, horizon: int) -> dict[str, np.ndarray]:
    """All inputs precede the forecast origin. Rows are individual SKU/unit series."""
    n = history.shape[1]
    weeks = history[:, -56:].reshape(len(history), 8, 7).sum(axis=2)[:, ::-1]
    weights = 0.72 ** np.arange(8)
    level = weeks @ weights / weights.sum()
    # Test a conservative spike cap, not a claim to identify customer/project orders.
    cap = np.array([2 * np.median(w[w > 0]) if np.any(w > 0) else 0 for w in weeks])
    robust = np.minimum(weeks, cap[:, None]) @ weights / weights.sum()
    result = {
        "sultan_mean59": history[:, -59:].mean(axis=1) * horizon,
        "mean28": history[:, -28:].mean(axis=1) * horizon,
        "weekly_ewma": level / 7 * horizon,
        "robust_ewma": robust / 7 * horizon,
    }
    seasonal = result["weekly_ewma"].copy()
    blended = seasonal.copy()
    if n >= 364:
        seasonal = history[:, n - 364:n - 364 + horizon].sum(axis=1)
    if n >= 420:
        previous = history[:, n - 420:n - 364].sum(axis=1)
        recent = history[:, -56:].sum(axis=1)
        ratio = np.divide(recent, previous, out=np.ones_like(recent), where=previous > 0)
        ratio = np.clip(ratio, 0.5, 2.0)
        blended = 0.5 * result["weekly_ewma"] + 0.5 * seasonal * ratio
    result["seasonal364"] = seasonal
    result["seasonal_blend"] = blended
    return result


def order_quantity(demand, free_stock, transit, multiple=1):
    """Illustrative replenishment arithmetic, not a production order engine."""
    if free_stock is None or free_stock < 0 or transit < 0 or multiple <= 0:
        raise ValueError("Missing/invalid stock or planning inputs")
    return math.ceil(max(0, demand - free_stock - transit) / multiple) * multiple


def availability_forecast(recorded_sales, available_days, horizon):
    """Synthetic demonstration; real exact availability is absent from the sources."""
    if available_days <= 0:
        raise ValueError("No observed availability to estimate a rate")
    return recorded_sales / available_days * horizon


def self_check():
    steady = np.ones((1, 500)) * 10
    for value in forecast(steady, 28).values():
        assert np.allclose(value, 280)
    for value in forecast(np.zeros((1, 500)), 7).values():
        assert np.allclose(value, 0)
    # The same call with later observations appended outside the history slice is identical.
    augmented = np.concatenate([steady, np.ones((1, 28)) * 999999], axis=1)
    for model, value in forecast(steady, 28).items():
        assert np.array_equal(value, forecast(augmented[:, :500], 28)[model])
    assert order_quantity(200, 23, 120, 6) == 60
    assert order_quantity(200, 23, 0, 6) == 180
    assert order_quantity(200, 500, 120, 6) == 0
    assert availability_forecast(140, 14, 28) == 280
    try:
        order_quantity(200, None, 120)
    except ValueError:
        pass
    else:
        raise AssertionError("Unknown stock must not become zero")


def load_supplier(folder, source_rows=None):
    path = next(folder.glob("Динамика*"))
    workbook = load_workbook(path, read_only=True, data_only=True)
    size = (END - START).days + 1
    series, names, issue_keys = {}, {}, set()
    counts = Counter()
    quantities = defaultdict(float)
    warehouses = set()
    dates = []
    for row_number, row in enumerate(workbook.active.iter_rows(min_row=2, values_only=True), start=2):
        if source_rows is not None:
            source_rows.append((row_number, row[:8]))
        stamp, _, document, sku, name, unit, warehouse, qty = row[:8]
        if stamp is None or sku is None:
            continue
        counts["source_rows"] += 1
        day = stamp.date() if isinstance(stamp, datetime) else datetime.strptime(str(stamp), "%d.%m.%Y %H:%M:%S").date()
        dates.append(day)
        key = (str(sku), str(unit))
        warehouses.add(str(warehouse))
        if day < START or day > END:
            counts["outside_experiment_dates"] += 1
            continue
        if not isinstance(qty, (int, float)) or not math.isfinite(qty):
            counts["missing_or_invalid_qty"] += 1
            issue_keys.add(key)
            continue
        if qty <= 0 or not str(document).startswith("Расходная накладная"):
            counts["ambiguous_or_non_outgoing_rows"] += 1
            quantities[f"excluded_signed_qty_{unit}"] += qty
            issue_keys.add(key)
            continue
        if key not in series:
            series[key] = np.zeros(size)
            names[key] = str(name)
        series[key][(day - START).days] += qty
        counts["included_positive_outgoing_rows"] += 1
        quantities[f"included_qty_{unit}"] += qty
    workbook.close()
    keys = sorted(series)
    matrix = np.array([series[key] for key in keys])
    assert np.isclose(matrix.sum(), sum(v for k, v in quantities.items() if k.startswith("included_")))
    metadata = {
        "source": str(path.relative_to(ROOT)), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "sheet": "Лист_1", "columns": "A:H", "counts": dict(counts),
        "quantities_by_unit": dict(quantities), "sku_unit_series": len(keys),
        "sku_count": len({key[0] for key in keys}), "issue_series": len(issue_keys),
        "warehouses": sorted(warehouses), "source_min_date": min(dates).isoformat(),
        "source_max_date": max(dates).isoformat(),
    }
    print(json.dumps({"supplier": folder.name, **metadata["counts"]}, ensure_ascii=True), flush=True)
    return keys, names, matrix, issue_keys, metadata


def score(rows, model):
    actual = np.array([r["actual"] for r in rows])
    predicted = np.array([r[model] for r in rows])
    total = actual.sum()
    error = predicted - actual
    return {
        "n": len(rows), "sku_count": len({r["sku"] for r in rows}),
        "actual_qty": round(float(total), 3),
        "predicted_qty": round(float(predicted.sum()), 3),
        "wape_pct": round(float(np.abs(error).sum() / total * 100), 3) if total else None,
        "bias_pct": round(float(error.sum() / total * 100), 3) if total else None,
        "underforecast_pct": round(float(np.maximum(-error, 0).sum() / total * 100), 3) if total else None,
        "overforecast_pct": round(float(np.maximum(error, 0).sum() / total * 100), 3) if total else None,
        "zero_actual_pct": round(float((actual == 0).mean() * 100), 3),
    }


def write_csv(path, rows):
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def main():
    self_check()
    OUT.mkdir(parents=True, exist_ok=True)
    records, sources, coverage = [], [], []
    origins = [date(2025, month, 1) for month in range(7, 13)] + [date(2026, month, 1) for month in range(1, 9)]
    for folder in sorted((ROOT / "docs/data").iterdir()):
        keys, names, matrix, issue_keys, metadata = load_supplier(folder)
        sources.append({"supplier": folder.name, **metadata})
        first_sale = np.argmax(matrix > 0, axis=1)
        for origin in origins:
            offset = (origin - START).days
            # Inclusion depends only on history, not on future sales or survival.
            eligible = first_sale <= offset - 90
            indexes = np.flatnonzero(eligible)
            history = matrix[eligible, :offset]
            active_weeks = (history[:, -56:].reshape(len(history), 8, 7).sum(axis=2) > 0).sum(axis=1)
            for horizon in (7, 28):
                forecasts = forecast(history, horizon)
                actual = matrix[eligible, offset:offset + horizon].sum(axis=1)
                for unit in sorted({key[1] for key in keys}):
                    unit_mask = np.array([key[1] == unit for key in keys])
                    total = matrix[unit_mask, offset:offset + horizon].sum()
                    covered = matrix[unit_mask & eligible, offset:offset + horizon].sum()
                    coverage.append({"supplier": folder.name, "origin": origin.isoformat(), "horizon": horizon,
                                     "unit": unit, "all_observed_qty": float(total), "eligible_qty": float(covered),
                                     "eligible_series": int(np.sum(unit_mask & eligible))})
                for j, index in enumerate(indexes):
                    sku, unit = keys[index]
                    records.append({
                        "supplier": folder.name, "unit": unit, "sku": sku, "name": names[keys[index]],
                        "origin": origin.isoformat(), "horizon": horizon,
                        "phase": "validation_2025" if origin.year == 2025 else "holdout_2026",
                        "segment": "regular" if active_weeks[j] >= 4 else "intermittent",
                        "has_ambiguous_rows_any_experiment_date": keys[index] in issue_keys,
                        "actual": float(actual[j]), **{m: float(v[j]) for m, v in forecasts.items()},
                    })
    selections = {}
    validation = defaultdict(list)
    for row in records:
        if row["phase"] == "validation_2025":
            validation[(row["supplier"], row["unit"], row["horizon"])].append(row)
    for key, rows in validation.items():
        selections[key] = min(SHORT_MODELS, key=lambda m: score(rows, m)["wape_pct"])
    for row in records:
        row["selected_2025"] = row[selections[(row["supplier"], row["unit"], row["horizon"])]]
    grouped = defaultdict(list)
    monthly = defaultdict(list)
    for row in records:
        base = (row["phase"], row["supplier"], row["unit"], row["horizon"])
        grouped[(*base, "all")].append(row)
        grouped[(*base, row["segment"])].append(row)
        if not row["has_ambiguous_rows_any_experiment_date"]:
            grouped[(*base, "no_ambiguous_rows_sensitivity")].append(row)
        monthly[(row["origin"], row["supplier"], row["unit"], row["horizon"])].append(row)
    metrics = []
    for (phase, supplier, unit, horizon, segment), rows in grouped.items():
        for model in MODELS + ["selected_2025"]:
            metrics.append({"phase": phase, "supplier": supplier, "unit": unit, "horizon": horizon,
                            "segment": segment, "model": model, **score(rows, model)})
    monthly_metrics = []
    for (origin, supplier, unit, horizon), rows in monthly.items():
        for model in MODELS + ["selected_2025"]:
            monthly_metrics.append({"origin": origin, "supplier": supplier, "unit": unit, "horizon": horizon,
                                    "model": model, **score(rows, model)})
    # Deliberately synthetic: demand is known, unlike censored real sales.
    normal = np.ones((1, 500)) * 10
    spike = normal.copy()
    spike[0, -4] += 1000
    spike_results = {m: {"normal": float(forecast(normal, 28)[m][0]),
                         "one_off_plus_1000": float(forecast(spike, 28)[m][0])} for m in MODELS}
    scenarios = {
        "label": "SYNTHETIC mechanics tests; not measured business outcomes",
        "spike": spike_results,
        "known_stockout": {"true_daily_demand": 10, "calendar_days": 28, "observed_in_stock_days": 14,
                           "recorded_sales": 140, "raw_next_28d": 140,
                           "availability_adjusted_next_28d": availability_forecast(140, 14, 28)},
        "transit_and_multiple": {"assumed_demand": 200, "assumed_free_stock": 23, "assumed_multiple": 6,
                                 "order_with_transit_120": order_quantity(200, 23, 120, 6),
                                 "order_with_transit_0": order_quantity(200, 23, 0, 6)},
    }
    assert scenarios["spike"]["robust_ewma"]["one_off_plus_1000"] < scenarios["spike"]["sultan_mean59"]["one_off_plus_1000"]
    summary = {
        "experiment_start": START.isoformat(), "experiment_end": END.isoformat(),
        "origin_dates": [o.isoformat() for o in origins], "horizons": [7, 28],
        "source_files": sources, "forecast_rows": len(records),
        "models_selected_using_2025_only": [{"supplier": k[0], "unit": k[1], "horizon": k[2], "model": v}
                                             for k, v in selections.items()],
        "metrics": metrics, "synthetic_scenarios": scenarios,
    }
    (OUT / "results.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    write_csv(OUT / "metrics.csv", metrics)
    write_csv(OUT / "monthly_metrics.csv", monthly_metrics)
    write_csv(OUT / "coverage.csv", coverage)
    write_csv(OUT / "predictions.csv", records)
    print(json.dumps({"forecast_rows": len(records), "selected_models": summary["models_selected_using_2025_only"]},
                     ensure_ascii=True), flush=True)


if __name__ == "__main__":
    main()
